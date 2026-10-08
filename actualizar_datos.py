import openpyxl, re, json, unicodedata, sys, os, time
from collections import Counter

sys.stdout.reconfigure(encoding='utf-8')

# File paths
retail_path = r'C:\Users\Lenovo\Documents\Mi dashboard\Dashboard HTML RETAIL .xlsx'
prepago_path = r'C:\Users\Lenovo\Documents\Mi dashboard\Dashboard HTML PREPAGO.xlsx'
p2_path = r'C:\Users\Lenovo\Documents\Mi dashboard\RETAIL P2.xlsx'

index_path = r'C:\Users\Lenovo\Documents\Mi dashboard\index.html'
dash_path = r'C:\Users\Lenovo\Documents\Mi dashboard\dashboard.html'
json_path = r'C:\Users\Lenovo\Documents\Mi dashboard\data_retail.json'
json_p2_path = r'C:\Users\Lenovo\Documents\Mi dashboard\data_retail_p2.json'
web_repo_path = r'C:\Users\Lenovo\Documents\salesland-dashboard-web\index.html'
web_repo_dir = r'C:\Users\Lenovo\Documents\salesland-dashboard-web'

def strip_accents(text):
    text = unicodedata.normalize('NFD', str(text or ''))
    text = ''.join(c for c in text if unicodedata.category(c) != 'Mn')
    return text.lower().strip()

def clean_name(name):
    s = strip_accents(name)
    s = re.sub(r'[^a-z0-9\s]', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip()
    return s

def clean_price(val):
    if not val: return 0
    s = str(val).replace('₡', '').replace('.', '').replace(',', '').strip()
    m = re.search(r'(\d+)', s)
    if m:
        try:
            num = int(m.group(1))
            if num > 1000000: return 0
            return num
        except: return 0
    return 0

def clean_str_val(val):
    if val is None: return ''
    s = str(val).strip()
    if s.endswith('.0'): s = s[:-2]
    if s in ['-', 'N/A', 'NA', 'None', 'null', 'nan']: return ''
    return s

meses_validos = {'enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre'}
n_to_m = {'01':'Enero','02':'Febrero','03':'Marzo','04':'Abril','05':'Mayo','06':'Junio','07':'Julio','08':'Agosto','09':'Septiembre','10':'Octubre','11':'Noviembre','12':'Diciembre'}

def normalize_retail_service(raw_val):
    s = str(raw_val or '').strip().lower()
    if 'fonatel' in s or 'ifi' in s:
        return 'IFI'
    elif 'portabilidad' in s:
        return 'Portabilidad Postpago'
    elif 'migraci' in s:
        return 'Migración Pospago'
    elif 'gpon' in s or 'fibra' in s:
        return 'Gpon'
    elif 'dth' in s:
        return 'DTH'
    elif 'postpago' in s or 'pospago' in s:
        return 'Pospago'
    return 'Pospago'

def normalize_provincia(p_raw):
    p = str(p_raw or '').strip()
    p_lower = strip_accents(p)
    if 'san jose' in p_lower: return 'San José'
    if 'alajuela' in p_lower: return 'Alajuela'
    if 'cartago' in p_lower: return 'Cartago'
    if 'heredia' in p_lower: return 'Heredia'
    if 'guanacaste' in p_lower or 'nicoya' in p_lower: return 'Guanacaste'
    if 'puntarenas' in p_lower: return 'Puntarenas'
    if 'limon' in p_lower: return 'Limón'
    return p if p else 'Sin Especificar'

EXPLICIT_MAP = {
    "Jesus Eduardo Miranda Piñar": "Jesus Miranda Piñar",
    "Sochil Pamela Ovares Chacón": "Zoshil Pamela Ovares Chacon",
    "Yorleny De Los Angeles Vargas Rojas": "Yorleni De Los Angeles Vargas Rojas",
    "Carlos Flores Sánchez": "Carlos Enrique Flores Sanchez",
    "Victoria Ruiz Villalobos": "Marianne Victoria Ruiz Villalobos",
    "Seidy Cruz Céspedes": "Seidy Delfina Cruz Céspedes",
    "Seidy Delfina Cruz Cespedes": "Seidy Delfina Cruz Céspedes",
    "Andy Solano Mendez": "Andy Gabriel Solano Mendez",
    "Cristofer Alonso Vargas": "Cristopher Alonso Vargas",
    "Noilyn Alvarez Cascante": "Noylin Erlith Álvarez Cascante",
    "Sherryl Segura Solano": "Sheryll Segura Solano",
    "Yorlene De Los Ángeles García": "Yorlene De Los Angeles Garcia",
    "Angie Maria Jiménez Castillo": "Angie Maria Jimenez Castillo",
    "Eliette Nicole Rodríguez Arguedas": "Eliette Nicole Rodriguez Arguedas",
    "Gendry Alvarez Obando": "Gendry Álvarez Obando",
    "Kemsy Krisey Briones Jiménez": "Kemsy Krisey Briones Jimenez",
    "Pablo Alonso Chavarría Mendiola": "Pablo Alonso Chavarria Mendiola",
    "Daniela Dayana Zuñiga Rodriguez": "Daniela Dayana Zúñiga Rodriguez",
    "Keisy Daniela Chavarria Arroyo": "Keisy Daniela Chavarría Arroyo",
    "Yendry Paola Vargas Sandi": "Yendry Paola Vargas Sandí",
    "Yoselin Maria Duran Montero": "Yoselin María Duran Montero",
    "Evelyn Patricia Gonzalez Arrieta": "Evelyn Patricia González Arrieta",
    "Kimberly Fuentes González": "Kimberly Fuentes González",
    "Nicole María Seas Correa": "Nicole María Seas Correa",
    "Maria Daniela Molina Castro": "María Daniela Molina Castro",
    "María Gabriela Rojas Carvajal": "Maria Gabriela Rojas Carvajal",
    "Stephanie Davila Vasquez": "Stephanie Dávila Vásquez",
    "Jennifer Maria Rodriguez Calderon": "Jennifer Maria Rodríguez Calderon",
    "Miriam Maria Porras Mora": "Míriam Maria Porras Mora",
    "Ericka Alexandra Hernandez Romero": "Ericka Hernández Romero",
    "Deybel Pricsilla Pena Lopez": "Deybell Priscilla Peña",
    "Lisseth Anabell Campos Colimdres": "Lisseth Campos Colindres"
}

# ==========================================
# 1. PROCESSING RETAIL P1
# ==========================================
print("======================================================================")
print("  SALESLAND | PROCESANDO DATOS RETAIL P1...")
print("======================================================================")
wb_retail = openpyxl.load_workbook(retail_path, read_only=True)
ws_retail = wb_retail['POSPAGO P1']

retail_records = []
retail_sellers = set()

for r in ws_retail.iter_rows(min_row=2, values_only=True):
    if not any(c is not None and str(c).strip() != '' for c in r[:16]): continue
    f = r[0]
    fecha_str = f.strftime('%Y-%m-%d') if hasattr(f, 'strftime') else str(f)[:10] if f else ''
    
    anio = r[1]
    try: anio = int(anio)
    except:
        if fecha_str and len(fecha_str) >= 4:
            try: anio = int(fecha_str[:4])
            except: anio = 2026
        else: anio = 2026
            
    mes = str(r[2]).strip() if r[2] else ''
    if mes.lower() not in meses_validos:
        if fecha_str and len(fecha_str) >= 7:
            m_num = fecha_str[5:7]
            mes = n_to_m.get(m_num, 'Sin Mes')
        else: mes = 'Sin Mes'
    else: mes = mes.capitalize()
        
    contrato = clean_str_val(r[4])
    id_tramite = clean_str_val(r[3])
    cliente = str(r[5]).strip() if r[5] else 'Cliente Final'
    cedula = clean_str_val(r[6])
    plan = str(r[7]).strip() if r[7] else 'Sin Plan'
    precio = clean_price(r[8])
    
    raw_canal = str(r[9] or '').strip().lower()
    if 'cadena' in raw_canal: canal = 'Cadenas'
    elif 'multi' in raw_canal: canal = 'Multimarca'
    elif 'avicam' in raw_canal: canal = 'Avicam'
    else: canal = 'Otros'
        
    pdv = str(r[10]).strip() if r[10] else 'General'
    norm_tipo = normalize_retail_service(r[11])
    
    raw_est = str(r[12] or '').strip().upper()
    estado = 'VENTA' if raw_est == 'VENTA' else 'RECHAZADA'
    
    vendedor = str(r[13]).strip() if r[13] else 'Sin Asignar'
    if vendedor in ['-', '', 'NA', 'N/A']: vendedor = 'Sin Asignar'
    retail_sellers.add(vendedor)
    
    supervisor = str(r[14]).strip() if r[14] else 'Sin Supervisor'
    if supervisor in ['-', '', 'NA', 'N/A']: supervisor = 'Sin Supervisor'
    
    telefono = clean_str_val(r[15]) or clean_str_val(r[16])
    
    raw_don = str(r[17] or '').strip().lower()
    if 'liberty' in raw_don: donante = 'Liberty'
    elif 'kolbi' in raw_don or 'kölbi' in raw_don: donante = 'Kölbi'
    else: donante = 'N/A'
        
    provincia = normalize_provincia(r[21] if len(r) > 21 else '')
    canton = str(r[22]).strip() if len(r) > 22 and r[22] else ''
    
    retail_records.append([fecha_str, anio, mes, cliente, cedula, plan, precio, canal, pdv, norm_tipo, estado, vendedor, supervisor, donante, provincia, canton, contrato, id_tramite, telefono])

wb_retail.close()
print(f"Pospago P1 extraído: {len(retail_records):,} registros")

retail_cleaned_map = {clean_name(s): s for s in retail_sellers}

def resolve_prepago_seller(asesor_raw, ref_sellers, ref_clean_map):
    raw = str(asesor_raw or '').strip()
    if not raw or raw in ['-', 'N/A', 'NA']: return 'Sin Asignar'
    if raw in EXPLICIT_MAP: return EXPLICIT_MAP[raw]
    if raw in ref_sellers: return raw
    cleaned = clean_name(raw)
    if cleaned in ref_clean_map: return ref_clean_map[cleaned]
    return raw

wb_prepago = openpyxl.load_workbook(prepago_path, read_only=True)
ws_prepago = wb_prepago['PREPAGO P1']

prepago_records = []

for r in ws_prepago.iter_rows(min_row=2, values_only=True):
    if not any(c is not None and str(c).strip() != '' for c in r[:10]): continue
    f = r[0]
    fecha_str = f.strftime('%Y-%m-%d') if hasattr(f, 'strftime') else str(f)[:10] if f else ''
    
    anio = r[1]
    try: anio = int(anio)
    except:
        if fecha_str and len(fecha_str) >= 4:
            try: anio = int(fecha_str[:4])
            except: anio = 2026
        else: anio = 2026
            
    mes = str(r[2]).strip() if r[2] else ''
    if mes.lower() not in meses_validos:
        if fecha_str and len(fecha_str) >= 7:
            m_num = fecha_str[5:7]
            mes = n_to_m.get(m_num, 'Sin Mes')
        else: mes = 'Sin Mes'
    else: mes = mes.capitalize()
    
    asesor_raw = r[3]
    vendedor = resolve_prepago_seller(asesor_raw, retail_sellers, retail_cleaned_map)
    
    cliente = str(r[4]).strip() if r[4] else 'Cliente Final'
    cedula = clean_str_val(r[5])
    precio = clean_price(r[7])
    
    raw_serv = str(r[9] or '').strip().lower()
    if 'porta' in raw_serv: norm_tipo = 'Portabilidad Prepago'
    else: norm_tipo = 'Prepago'
    plan = norm_tipo
    
    raw_canal = str(r[16] or '').strip().lower() if len(r) > 16 else ''
    if 'cadena' in raw_canal: canal = 'Cadenas'
    elif 'multi' in raw_canal: canal = 'Multimarca'
    elif 'avicam' in raw_canal: canal = 'Avicam'
    else: canal = 'Otros'
    
    pdv = str(r[17]).strip() if len(r) > 17 and r[17] else 'General'
    estado = 'VENTA'
    
    supervisor = str(r[14]).strip() if len(r) > 14 and r[14] else 'Sin Supervisor'
    if supervisor in ['-', '', 'NA', 'N/A']: supervisor = 'Sin Supervisor'
    
    raw_don = str(r[11] or '').strip().lower() if len(r) > 11 else ''
    if 'liberty' in raw_don: donante = 'Liberty'
    elif 'kolbi' in raw_don or 'kölbi' in raw_don: donante = 'Kölbi'
    else: donante = 'N/A'
    
    provincia = normalize_provincia(r[18] if len(r) > 18 else '')
    canton = str(r[19]).strip() if len(r) > 19 and r[19] else ''
    
    contrato = 'Prepago'
    id_tramite = clean_str_val(r[8]) if len(r) > 8 else ''
    telefono = clean_str_val(r[6]) if len(r) > 6 else ''
    if not telefono and len(r) > 10: telefono = clean_str_val(r[10])
    
    prepago_records.append([fecha_str, anio, mes, cliente, cedula, plan, precio, canal, pdv, norm_tipo, estado, vendedor, supervisor, donante, provincia, canton, contrato, id_tramite, telefono])

wb_prepago.close()
print(f"Prepago P1 extraído: {len(prepago_records):,} registros")

combined_records_p1 = retail_records + prepago_records
print(f"TOTAL RETAIL P1: {len(combined_records_p1):,} registros")

with open(json_path, 'w', encoding='utf-8') as f:
    json.dump(combined_records_p1, f, ensure_ascii=False)
print("Guardado data_retail.json con éxito!")

# ==========================================
# 2. PROCESSING RETAIL P2
# ==========================================
combined_records_p2 = []
if os.path.exists(p2_path):
    print("\n======================================================================")
    print("  SALESLAND | PROCESANDO DATOS RETAIL P2...")
    print("======================================================================")
    wb_p2 = openpyxl.load_workbook(p2_path, read_only=True)
    
    ws_pos_p2 = wb_p2['POSPAGO P2']
    pospago_p2_records = []
    p2_sellers = set()
    
    for r in ws_pos_p2.iter_rows(min_row=2, values_only=True):
        if not any(c is not None and str(c).strip() != '' for c in r[:16]): continue
        f = r[0]
        fecha_str = f.strftime('%Y-%m-%d') if hasattr(f, 'strftime') else str(f)[:10] if f else ''
        
        anio = r[1]
        try: anio = int(anio)
        except:
            if fecha_str and len(fecha_str) >= 4:
                try: anio = int(fecha_str[:4])
                except: anio = 2026
            else: anio = 2026
                
        mes = str(r[2]).strip() if r[2] else ''
        if mes.lower() not in meses_validos:
            if fecha_str and len(fecha_str) >= 7:
                m_num = fecha_str[5:7]
                mes = n_to_m.get(m_num, 'Sin Mes')
            else: mes = 'Sin Mes'
        else: mes = mes.capitalize()
            
        contrato = clean_str_val(r[4])
        id_tramite = clean_str_val(r[3])
        cliente = str(r[5]).strip() if r[5] else 'Cliente Final'
        cedula = clean_str_val(r[6])
        plan = str(r[7]).strip() if r[7] else 'Sin Plan'
        precio = clean_price(r[8])
        
        raw_canal = str(r[9] or '').strip().lower()
        if 'cadena' in raw_canal: canal = 'Cadenas'
        elif 'multi' in raw_canal: canal = 'Multimarca'
        elif 'avicam' in raw_canal: canal = 'Avicam'
        else: canal = 'Otros'
            
        pdv = str(r[10]).strip() if r[10] else 'General'
        norm_tipo = normalize_retail_service(r[11])
        
        raw_est = str(r[12] or '').strip().upper()
        estado = 'VENTA' if raw_est == 'VENTA' else 'RECHAZADA'
        
        vendedor = str(r[13]).strip() if r[13] else 'Sin Asignar'
        if vendedor in ['-', '', 'NA', 'N/A']: vendedor = 'Sin Asignar'
        p2_sellers.add(vendedor)
        
        supervisor = str(r[14]).strip() if r[14] else 'Sin Supervisor'
        if supervisor in ['-', '', 'NA', 'N/A']: supervisor = 'Sin Supervisor'
        
        telefono = clean_str_val(r[15]) or clean_str_val(r[16])
        
        raw_don = str(r[17] or '').strip().lower()
        if 'liberty' in raw_don: donante = 'Liberty'
        elif 'kolbi' in raw_don or 'kölbi' in raw_don: donante = 'Kölbi'
        else: donante = 'N/A'
            
        provincia = normalize_provincia(r[21] if len(r) > 21 else '')
        canton = str(r[22]).strip() if len(r) > 22 and r[22] else ''
        
        pospago_p2_records.append([fecha_str, anio, mes, cliente, cedula, plan, precio, canal, pdv, norm_tipo, estado, vendedor, supervisor, donante, provincia, canton, contrato, id_tramite, telefono])
    
    print(f"Pospago P2 extraído: {len(pospago_p2_records):,} registros")
    p2_cleaned_map = {clean_name(s): s for s in p2_sellers}
    
    ws_prep_p2 = wb_p2['PREPAGO P2']
    prepago_p2_records = []
    
    for r in ws_prep_p2.iter_rows(min_row=2, values_only=True):
        if not any(c is not None and str(c).strip() != '' for c in r[:10]): continue
        f = r[0]
        fecha_str = f.strftime('%Y-%m-%d') if hasattr(f, 'strftime') else str(f)[:10] if f else ''
        
        anio = r[1]
        try: anio = int(anio)
        except:
            if fecha_str and len(fecha_str) >= 4:
                try: anio = int(fecha_str[:4])
                except: anio = 2026
            else: anio = 2026
                
        mes = str(r[2]).strip() if r[2] else ''
        if mes.lower() not in meses_validos:
            if fecha_str and len(fecha_str) >= 7:
                m_num = fecha_str[5:7]
                mes = n_to_m.get(m_num, 'Sin Mes')
            else: mes = 'Sin Mes'
        else: mes = mes.capitalize()
        
        asesor_raw = r[3]
        vendedor = resolve_prepago_seller(asesor_raw, p2_sellers, p2_cleaned_map)
        
        cliente = str(r[4]).strip() if r[4] else 'Cliente Final'
        cedula = clean_str_val(r[5])
        precio = clean_price(r[7])
        
        raw_serv = str(r[9] or '').strip().lower()
        if 'porta' in raw_serv: norm_tipo = 'Portabilidad Prepago'
        else: norm_tipo = 'Prepago'
        plan = norm_tipo
        
        raw_canal = str(r[16] or '').strip().lower() if len(r) > 16 else ''
        if 'cadena' in raw_canal: canal = 'Cadenas'
        elif 'multi' in raw_canal: canal = 'Multimarca'
        elif 'avicam' in raw_canal: canal = 'Avicam'
        else: canal = 'Otros'
        
        pdv = str(r[17]).strip() if len(r) > 17 and r[17] else 'General'
        estado = 'VENTA'
        
        supervisor = str(r[14]).strip() if len(r) > 14 and r[14] else 'Sin Supervisor'
        if supervisor in ['-', '', 'NA', 'N/A']: supervisor = 'Sin Supervisor'
        
        raw_don = str(r[11] or '').strip().lower() if len(r) > 11 else ''
        if 'liberty' in raw_don: donante = 'Liberty'
        elif 'kolbi' in raw_don or 'kölbi' in raw_don: donante = 'Kölbi'
        else: donante = 'N/A'
        
        provincia = normalize_provincia(r[18] if len(r) > 18 else '')
        canton = str(r[19]).strip() if len(r) > 19 and r[19] else ''
        
        contrato = 'Prepago'
        id_tramite = clean_str_val(r[8]) if len(r) > 8 else ''
        telefono = clean_str_val(r[6]) if len(r) > 6 else ''
        if not telefono and len(r) > 10: telefono = clean_str_val(r[10])
        
        prepago_p2_records.append([fecha_str, anio, mes, cliente, cedula, plan, precio, canal, pdv, norm_tipo, estado, vendedor, supervisor, donante, provincia, canton, contrato, id_tramite, telefono])
    
    wb_p2.close()
    print(f"Prepago P2 extraído: {len(prepago_p2_records):,} registros")
    
    combined_records_p2 = pospago_p2_records + prepago_p2_records
    print(f"TOTAL RETAIL P2: {len(combined_records_p2):,} registros")
    
    with open(json_p2_path, 'w', encoding='utf-8') as f:
        json.dump(combined_records_p2, f, ensure_ascii=False)
    print("Guardado data_retail_p2.json con éxito!")

# ==========================================
# 3. EMBEDDING IN HTML FILES & OFFICIAL TARGETS
# ==========================================
print("\n======================================================================")
print("  EMBEBIENDO DATASETS Y METAS OFICIALES (P1 & P2)...")
print("======================================================================")
json_dataset_str_p1 = json.dumps(combined_records_p1, ensure_ascii=False)
json_dataset_str_p2 = json.dumps(combined_records_p2, ensure_ascii=False) if combined_records_p2 else "[]"

def update_dataset_in_html(file_path):
    with open(file_path, 'r', encoding='utf-8') as f:
        html = f.read()

    # 1. Update comment
    p2_count_str = f" y {len(combined_records_p2):,} registros (Retail P2)" if combined_records_p2 else ""
    html = re.sub(
        r'//\s*Embedded dataset of[^\n]+',
        f'// Embedded dataset of {len(combined_records_p1):,} records (Retail P1){p2_count_str}',
        html
    )

    # 2. Replace window.RAW_DATASET
    raw_dataset_regex = re.compile(r'window\.RAW_DATASET\s*=\s*\[\[[\s\S]*?\]\];', re.MULTILINE)
    if raw_dataset_regex.search(html):
        html = raw_dataset_regex.sub(f'window.RAW_DATASET = {json_dataset_str_p1};', html)
    else:
        print(f"ERROR: Could not find window.RAW_DATASET in {file_path}")
        return False

    # 3. Replace or insert window.RAW_DATASET_P2
    if 'window.RAW_DATASET_P2' in html:
        raw_p2_regex = re.compile(r'window\.RAW_DATASET_P2\s*=\s*\[[\s\S]*?\];', re.MULTILINE)
        html = raw_p2_regex.sub(f'window.RAW_DATASET_P2 = {json_dataset_str_p2};', html)
    else:
        target_marker = f'window.RAW_DATASET = {json_dataset_str_p1};'
        insert_code = f'\n    window.RAW_DATASET_P2 = {json_dataset_str_p2};'
        html = html.replace(target_marker, target_marker + insert_code, 1)

    # 4. Official Targets Definition (Retail P1: 7,806 | Retail P2: 6,914)
    targets_code_block = """    // Service Targets System by Month & Project (Official Quotas)
    const DEFAULT_TARGETS_P1 = {
      'Octubre': {
        'Prepago': 5450,
        'Portabilidad Prepago': 1000,
        'Pospago': 1296,
        'Gpon': 36,
        'DTH': 12,
        'IFI': 12
      },
      'Septiembre': {
        'Prepago': 5450,
        'Portabilidad Prepago': 1000,
        'Pospago': 1296,
        'Gpon': 36,
        'DTH': 24,
        'IFI': 0
      }
    };

    const DEFAULT_TARGETS_P2 = {
      'Octubre': {
        'Prepago': 4551,
        'Portabilidad Prepago': 1000,
        'Pospago': 1303,
        'Gpon': 36,
        'DTH': 12,
        'IFI': 12
      },
      'Septiembre': {
        'Prepago': 4551,
        'Portabilidad Prepago': 1000,
        'Pospago': 1303,
        'Gpon': 36,
        'DTH': 12,
        'IFI': 12
      }
    };

    // Metas Oficiales por Servicio Asignadas a Cada Vendedor (Ambos Proyectos Retail P1 y Retail P2)
    // Portabilidad Postpago: 10 | Pospago Línea Nueva: 25 | Prepago: 120 | Portabilidad Prepago: 15 | Gpon: 1 | DTH: 1 | IFI: 1
    // Total Meta Asignada por Vendedor: 173 instalaciones
    const DEFAULT_SELLER_TARGETS = {
      'Octubre': {
        'Prepago': 120,
        'Portabilidad Prepago': 15,
        'Pospago': 25,
        'Portabilidad Postpago': 10,
        'Gpon': 1,
        'DTH': 1,
        'IFI': 1
      },
      'Septiembre': {
        'Prepago': 120,
        'Portabilidad Prepago': 15,
        'Pospago': 25,
        'Portabilidad Postpago': 10,
        'Gpon': 1,
        'DTH': 1,
        'IFI': 1
      }
    };
    const DEFAULT_SUPERVISOR_TARGETS = DEFAULT_SELLER_TARGETS;"""

    # Replace existing DEFAULT_TARGETS_P1/P2 and SELLER targets block cleanly
    target_block_regex = re.compile(r'// Service Targets System[\s\S]*?(?=function getNormalizedTargetMonth)', re.MULTILINE)
    if target_block_regex.search(html):
        html = target_block_regex.sub(targets_code_block + '\n\n    ', html)
    else:
        fallback_regex = re.compile(r'const DEFAULT_TARGETS_P1\s*=\s*\{[\s\S]*?\};', re.MULTILINE)
        if fallback_regex.search(html):
            html = fallback_regex.sub(targets_code_block + '\n\n    ', html)

    # 5. Ensure getServiceTargets checks DEFAULT_TARGETS_P1 and DEFAULT_TARGETS_P2 with localStorage protection
    get_targets_fn = """    function getServiceTargets(month) {
      const proj = (typeof window !== 'undefined' && window.App && window.App.currentProject) ? window.App.currentProject : (localStorage.getItem('salesland_current_project') || 'P1');
      const normMonth = getNormalizedTargetMonth(month);

      // Version check to ensure official corporate quotas take precedence over outdated test storage
      const TARGETS_VERSION = 'v3_official_quotas_p1_p2';
      try {
        if (localStorage.getItem('salesland_targets_version') !== TARGETS_VERSION) {
          localStorage.removeItem('salesland_service_targets_P1_Octubre');
          localStorage.removeItem('salesland_service_targets_P1_Septiembre');
          localStorage.removeItem('salesland_service_targets_P2_Octubre');
          localStorage.removeItem('salesland_service_targets_P2_Septiembre');
          localStorage.setItem('salesland_targets_version', TARGETS_VERSION);
        }
        const key = 'salesland_service_targets_' + proj + '_' + normMonth;
        const saved = localStorage.getItem(key);
        if (saved) return JSON.parse(saved);
      } catch(e) {}

      if (proj === 'P1') {
        if (DEFAULT_TARGETS_P1[normMonth]) {
          return { ...DEFAULT_TARGETS_P1[normMonth] };
        }
        return { ...DEFAULT_TARGETS_P1['Octubre'] };
      }

      if (proj === 'P2') {
        if (DEFAULT_TARGETS_P2[normMonth]) {
          return { ...DEFAULT_TARGETS_P2[normMonth] };
        }
        return { ...DEFAULT_TARGETS_P2['Octubre'] };
      }

      return {
        'Prepago': 4551,
        'Portabilidad Prepago': 1000,
        'Pospago': 1303,
        'Gpon': 36,
        'DTH': 12,
        'IFI': 12
      };
    }"""

    get_targets_regex = re.compile(r'function getServiceTargets\(month\)\s*\{[\s\S]*?(?=function saveServiceTargets)', re.MULTILINE)
    if get_targets_regex.search(html):
        html = get_targets_regex.sub(get_targets_fn + '\n\n    ', html)

    # 6. Ensure App.dataP2 is initialized from window.RAW_DATASET_P2
    old_init_regex = re.compile(r'App\.currentProject\s*=\s*localStorage\.getItem\([^\)]+\)\s*\|\|\s*\'P1\';[\s\S]*?App\.dataP1\s*=\s*window\.RAW_DATASET\s*\|\|\s*\[\];[\s\S]*?switchProject\(', re.MULTILINE)
    new_init_code = """App.currentProject = localStorage.getItem('salesland_current_project') || 'P1';
    App.dataP1 = window.RAW_DATASET || [];
    App.dataP2 = (window.RAW_DATASET_P2 && window.RAW_DATASET_P2.length > 0) ? window.RAW_DATASET_P2 : [];
    if (App.dataP2.length === 0) {
      try {
        const savedP2 = localStorage.getItem('salesland_dataset_p2');
        App.dataP2 = savedP2 ? JSON.parse(savedP2) : [];
      } catch(e) {
        App.dataP2 = [];
      }
    }

    function switchProject("""
    if old_init_regex.search(html):
        html = old_init_regex.sub(new_init_code, html)

    # 7. Ensure switchProject resets supervisor & dia and calls updateDayFilterDropdown
    switch_fn_code = """function switchProject(proj) {
      App.currentProject = proj;
      localStorage.setItem('salesland_current_project', proj);

      const btnP1 = document.getElementById('btn-proj-p1');
      const btnP2 = document.getElementById('btn-proj-p2');
      const brandTitle = document.getElementById('brand-project-title');
      const p2Card = document.getElementById('p2-waiting-card');
      const btnClearP2 = document.getElementById('btn-clear-p2-data');

      if (proj === 'P2') {
        if (btnP1) btnP1.classList.remove('active');
        if (btnP2) btnP2.classList.add('active');
        if (brandTitle) brandTitle.textContent = 'Retail P2';
        document.title = 'Salesland Retail P2 - Panel Gerencial de Rendimiento Retail';

        if (App.dataP2 && App.dataP2.length > 0) {
          App.data = App.dataP2;
          if (p2Card) p2Card.style.display = 'none';
          if (btnClearP2) btnClearP2.style.display = 'none';
        } else {
          App.data = [];
          if (p2Card) p2Card.style.display = 'block';
          if (btnClearP2) btnClearP2.style.display = 'none';
        }
      } else {
        // Project P1
        if (btnP1) btnP1.classList.add('active');
        if (btnP2) btnP2.classList.remove('active');
        if (brandTitle) brandTitle.textContent = 'Retail P1';
        document.title = 'Salesland Retail P1 - Panel Gerencial de Rendimiento Retail';
        if (p2Card) p2Card.style.display = 'none';
        App.data = App.dataP1;
      }

      if (App.filters) {
        App.filters.supervisor = 'all';
        App.filters.dia = 'all';
        const supSelect = document.getElementById('filter-supervisor');
        if (supSelect) supSelect.value = 'all';
        const diaSelect = document.getElementById('filter-dia');
        if (diaSelect) diaSelect.value = 'all';
      }
      updateDayFilterDropdown();
      updateSupervisorFilterDropdown();
      applyAllFilters();
    }"""
    switch_regex = re.compile(r'function switchProject\(proj\)\s*\{[\s\S]*?applyAllFilters\(\);\s*\}', re.MULTILINE)
    if switch_regex.search(html):
        html = switch_regex.sub(switch_fn_code, html)

    # 8. Enhance Targets Modal to show current Project badge
    modal_title_old = '<span>🎯</span> Metas Comerciales por Servicio'
    modal_title_new = '<span>🎯</span> Metas Comerciales por Servicio <span id="target-modal-project-badge" style="font-size:0.75rem; background:#0284c7; color:#fff; padding:2px 8px; border-radius:12px; margin-left:auto;">Retail P1</span>'
    if modal_title_old in html and 'target-modal-project-badge' not in html:
        html = html.replace(modal_title_old, modal_title_new, 1)

    load_modal_old = "function loadTargetsToModal(month) {\n        const m = month || (App.filters.mes && App.filters.mes !== 'all' ? App.filters.mes : 'Octubre');"
    load_modal_new = """function loadTargetsToModal(month) {
        const proj = (window.App && window.App.currentProject) ? window.App.currentProject : 'P1';
        const badge = document.getElementById('target-modal-project-badge');
        if (badge) {
          badge.textContent = proj === 'P2' ? 'Retail P2' : 'Retail P1';
          badge.style.background = proj === 'P2' ? '#f59e0b' : '#0284c7';
        }
        const m = month || (App.filters.mes && App.filters.mes !== 'all' ? App.filters.mes : 'Octubre');"""
    if load_modal_old in html:
        html = html.replace(load_modal_old, load_modal_new, 1)

    # 9. Style project switcher buttons cleanly
    btn_p1_active_css = """.btn-project-pill.active#btn-proj-p1 {
      background: #0284c7 !important;
      color: #ffffff !important;
      box-shadow: 0 2px 6px rgba(2, 132, 199, 0.4);
    }
    .btn-project-pill.active#btn-proj-p2 {
      background: #f59e0b !important;
      color: #ffffff !important;
      box-shadow: 0 2px 6px rgba(245, 158, 11, 0.4);
    }"""
    if '.btn-project-pill.active#btn-proj-p1' not in html:
        html = html.replace('.btn-project-pill.active {', btn_p1_active_css + '\n    .btn-project-pill.active {', 1)

    with open(file_path, 'w', encoding='utf-8') as f:
        f.write(html)
    print(f"[OK] Metas y configuración actualizadas en {file_path}")
    return True

ok1 = update_dataset_in_html(index_path)
ok2 = update_dataset_in_html(dash_path)
ok3 = False
if os.path.exists(web_repo_path):
    ok3 = update_dataset_in_html(web_repo_path)

# ==========================================
# 4. GITHUB PAGES DEPLOYMENT
# ==========================================
if os.path.exists(os.path.join(web_repo_dir, '.git')):
    try:
        print("\nSubiendo actualización de Metas a GitHub Pages...")
        import subprocess
        import datetime
        ahora = datetime.datetime.now().strftime('%d/%m/%Y %H:%M')
        p2_txt = f", P2: {len(combined_records_p2):,} reg" if combined_records_p2 else ""
        msg = f"Actualización de datos (P1: {len(combined_records_p1):,} reg{p2_txt}) - {ahora}"
        subprocess.run(['git', 'add', 'index.html'], cwd=web_repo_dir, check=False)
        subprocess.run(['git', 'commit', '-m', msg], cwd=web_repo_dir, check=False)
        subprocess.run(['git', '-c', 'credential.helper=', '-c', 'credential.helper=manager', 'push', 'origin', 'main'], cwd=web_repo_dir, check=False)
        print("[OK] Publicacion en GitHub Pages completada con éxito.")
    except Exception as e:
        print(f"Nota Git: {e}")

print("\n======================================================================")
print(" ¡METAS ASIGNADAS Y DESPLEGADAS CON ÉXITO!")
print(" P1: Prepago 5450, Porta Prepago 1000, Pospago 1296, Gpon 36, DTH 12, IFI 12 (Total: 7,806)")
print(" P2: Prepago 4551, Porta Prepago 1000, Pospago 1303, Gpon 36, DTH 12, IFI 12 (Total: 6,914)")
print("======================================================================")
